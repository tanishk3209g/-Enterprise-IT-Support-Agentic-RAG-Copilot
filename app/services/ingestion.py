from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        docs = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": str(path), "page": i + 1}))
        return docs
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
        return [Document(page_content=text, metadata={"source": str(path)})]
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))
